from pptx import Presentation


def extract_slide_text(pptx_bytes: bytes) -> list[dict]:
    """Pull per-slide title, body text runs, and speaker notes from a PPTX.
    Deterministic, no LLM involved. Returns one dict per slide, 1-indexed.
    """
    import io

    presentation = Presentation(io.BytesIO(pptx_bytes))
    slides = []

    for index, slide in enumerate(presentation.slides, start=1):
        title = ""
        body_lines: list[str] = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape == slide.shapes.title:
                title = text
            else:
                body_lines.append(text)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        slides.append(
            {
                "slide_number": index,
                "title": title,
                "body": "\n".join(body_lines),
                "notes": notes,
            }
        )

    return slides
